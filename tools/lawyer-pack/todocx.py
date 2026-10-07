import re, glob, os
from docx import Document
from docx.shared import Pt
os.chdir('/home/user/Claude-code-Paint-Trade-HQ/lawyer')
def add_runs(par, text):
    text = text.replace('`', ''); pos = 0
    for m in re.finditer(r'\*\*(.+?)\*\*|\*(.+?)\*', text):
        if m.start() > pos: par.add_run(text[pos:m.start()])
        if m.group(1): par.add_run(m.group(1)).bold = True
        else: par.add_run(m.group(2)).italic = True
        pos = m.end()
    if pos < len(text): par.add_run(text[pos:])
for f in sorted(glob.glob('*.md')):
    doc = Document(); st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
    lines = open(f, encoding='utf-8').read().split('\n'); i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:\-|]+\|$', lines[i+1]):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s:\-|]+\|$', lines[i]): rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows)); t.style = 'Table Grid'
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    p = t.cell(r, c).paragraphs[0]; add_runs(p, cell)
                    if r == 0:
                        for run in p.runs: run.bold = True
            doc.add_paragraph(); continue
        m = re.match(r'^(#{1,3})\s+(.*)', ln)
        if m: doc.add_heading(m.group(2).replace('`', ''), level=len(m.group(1))); i += 1; continue
        if ln.startswith('> '): p = doc.add_paragraph(); add_runs(p, ln[2:]); i += 1; continue
        m = re.match(r'^(\s*)[-*]\s+(.*)', ln)
        if m: add_runs(doc.add_paragraph(style='List Bullet'), m.group(2)); i += 1; continue
        m = re.match(r'^\s*\d+\.\s+(.*)', ln)
        if m: add_runs(doc.add_paragraph(style='List Number'), m.group(1)); i += 1; continue
        if ln.strip() == '': i += 1; continue
        add_runs(doc.add_paragraph(), ln); i += 1
    doc.save(f[:-3] + '.docx'); print('docx', f)
